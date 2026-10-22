"""Adds the team to the deck's title slide, writes each slide's script into its speaker notes,
and produces a printable scripts document.

Usage: python add_team.py <deck.pptx> <scripts.docx>
The deck is updated in place (via a temporary copy that must pass validation first).
"""
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

from scripts_content import MEMBER_ORDER, MEMBERS, SCRIPTS, SLIDE_TITLES

VALIDATE = "/home/shafreed/.claude/skills/synced/f7d7525f-fac9-4ce7-b48f-0dbec9262a24_7916b980-97b8-4f86-91a9-c6e63cf9e18d/pptx/scripts/office/validate.py"
WPM = 130
EMU = 914400


def check_timing():
    for n, (who, minutes, paras) in sorted(SCRIPTS.items()):
        words = sum(len(p.split()) for p in paras)
        if words > minutes * WPM * 1.25:
            print(f"warning: slide {n} script is {words} words for {minutes} min ({MEMBERS[who][0]})")


# ------------------------------------------------------------------ deck ---
def title_slide(x):
    """Replace the placeholder 'Submitted by' line with the five members on two lines."""
    names = [f"{MEMBERS[k][0]} ({MEMBERS[k][1]})" for k in MEMBER_ORDER]
    m = re.search(r"<a:p>(?:(?!</a:p>).)*Submitted by:(?:(?!</a:p>).)*</a:p>", x)
    assert m, "title slide placeholder not found"
    old_p = m.group(0)
    rpr = re.search(r"<a:rPr[^>]*>.*?</a:rPr>", old_p).group(0)
    ppr = re.search(r"<a:pPr[^>]*>.*?</a:pPr>", old_p).group(0)
    bold_rpr = rpr.replace('<a:rPr lang="en-US"', '<a:rPr lang="en-US" b="1"', 1)

    def p(text, r=rpr):
        return f"<a:p>{ppr}<a:r>{r}<a:t>{escape(text)}</a:t></a:r></a:p>"

    lines = p("Submitted by", bold_rpr) + p(" · ".join(names[:3])) + p(" · ".join(names[3:]))
    x = x.replace(old_p, lines)
    # Give the text box room for two more lines: y = 4.85 in, height = 2.25 in
    box = re.search(r'<a:off x="(\d+)" y="4846320"/><a:ext cx="(\d+)" cy="1280160"/>', x)
    assert box, "title slide text box not found"
    x = x.replace(box.group(0), f'<a:off x="{box.group(1)}" y="{int(4.85 * EMU)}"/>'
                                 f'<a:ext cx="{box.group(2)}" cy="{int(2.25 * EMU)}"/>')
    return x


def notes_xml_paragraphs(slide_no):
    who, minutes, paras = SCRIPTS[slide_no]
    name, roll = MEMBERS[who]
    head = f"Presenter: {name} ({roll}) · about {minutes:g} min"
    out = f'<a:p><a:r><a:rPr lang="en-US" b="1" dirty="0"/><a:t>{escape(head)}</a:t></a:r></a:p>'
    for t in paras:
        out += f'<a:p><a:r><a:rPr lang="en-US" dirty="0"/><a:t>{escape(t)}</a:t></a:r></a:p>'
    return out


def update_deck(deck):
    work = Path(tempfile.mkdtemp(prefix="deck_pkg_", dir=Path(__file__).parent))
    with zipfile.ZipFile(deck) as z:
        z.extractall(work)
    for link in [p for p in work.rglob("*") if p.is_symlink()]:
        link.unlink()

    s1 = work / "ppt" / "slides" / "slide1.xml"
    s1.write_text(title_slide(s1.read_text(encoding="utf8")), encoding="utf8")

    empty_body = re.compile(r'(<p:ph type="body" idx="1"/>.*?<a:lstStyle/>)(.*?)(</p:txBody>)', re.S)
    for n in SCRIPTS:
        rels = (work / "ppt" / "slides" / "_rels" / f"slide{n}.xml.rels").read_text(encoding="utf8")
        target = re.search(r'Target="\.\./notesSlides/(notesSlide\d+\.xml)"', rels).group(1)
        npath = work / "ppt" / "notesSlides" / target
        nx = npath.read_text(encoding="utf8")
        nx, count = empty_body.subn(lambda m: m.group(1) + notes_xml_paragraphs(n) + m.group(3), nx, count=1)
        assert count == 1, f"notes body not found for slide {n}"
        npath.write_text(nx, encoding="utf8")

    tmp = Path(deck).with_name("deck_candidate.pptx")
    if tmp.exists():
        tmp.unlink()
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        ct = work / "[Content_Types].xml"
        z.write(ct, "[Content_Types].xml")
        for f in sorted(work.rglob("*")):
            if f.is_file() and f != ct:
                z.write(f, f.relative_to(work).as_posix())
    shutil.rmtree(work)
    res = subprocess.run([sys.executable, VALIDATE, str(tmp), "--original", str(deck)], capture_output=True, text=True)
    print(res.stdout.strip().splitlines()[-1] if res.stdout.strip() else res.stderr)
    if res.returncode != 0:
        raise SystemExit(f"validation failed; deck left unchanged. Candidate: {tmp}")
    shutil.move(tmp, deck)


# --------------------------------------------------------------- scripts ---
def write_scripts_doc(path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(12)

    doc.add_heading("AlphaGo Zero: Presentation Scripts", level=0)
    doc.add_paragraph("CSE 311 Artificial Intelligence · Indian Institute of Information Technology Kottayam")
    doc.add_paragraph(f"Timings assume about {WPM} spoken words per minute. Each script is also in the speaker "
                      "notes of AlphaGoZero_Presentation.pptx, so presenter view shows it during the talk.")

    doc.add_heading("Running order", level=1)
    order, seen = [], {}
    for n in sorted(SCRIPTS):
        who = SCRIPTS[n][0]
        if order and order[-1][0] == who:
            order[-1][1].append(n)
        else:
            order.append((who, [n]))
        seen.setdefault(who, 0)
        seen[who] += SCRIPTS[n][1]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ["Part", "Presenter", "Slides", "Time"]):
        cell.text = text
    for k, (who, slides) in enumerate(order, 1):
        mins = sum(SCRIPTS[n][1] for n in slides)
        span = f"{slides[0]}–{slides[-1]}" if len(slides) > 1 else str(slides[0])
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = str(k), MEMBERS[who][0], span, f"{mins:g} min"
    total = sum(v[1] for v in SCRIPTS.values())
    doc.add_paragraph(f"Total speaking time: about {total:g} minutes.")

    doc.add_heading("Time per member", level=1)
    for who in MEMBER_ORDER:
        name, roll = MEMBERS[who]
        slides = [n for n in sorted(SCRIPTS) if SCRIPTS[n][0] == who]
        doc.add_paragraph(f"{name} ({roll}): slides {', '.join(map(str, slides))}, about {seen[who]:g} min",
                          style="List Bullet")

    for who in MEMBER_ORDER:
        name, roll = MEMBERS[who]
        doc.add_page_break()
        doc.add_heading(f"{name} ({roll})", level=1)
        slides = [n for n in sorted(SCRIPTS) if SCRIPTS[n][0] == who]
        doc.add_paragraph(f"Slides {', '.join(map(str, slides))} · about {seen[who]:g} minutes in total")
        for n in slides:
            _, minutes, paras = SCRIPTS[n]
            doc.add_heading(f"Slide {n}: {SLIDE_TITLES[n]}", level=2)
            meta = doc.add_paragraph(f"about {minutes:g} min")
            meta.runs[0].italic = True
            meta.runs[0].font.color.rgb = RGBColor(0x5A, 0x64, 0x72)
            for t in paras:
                doc.add_paragraph(t)
    # python-docx's default template omits the required w:percent on w:zoom
    from docx.oxml.ns import qn
    for zoom in doc.settings.element.iter(qn("w:zoom")):
        if zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")
    doc.save(path)


if __name__ == "__main__":
    deck, scripts_doc = sys.argv[1:3]
    check_timing()
    update_deck(deck)
    write_scripts_doc(scripts_doc)
    print("scripts written to", scripts_doc)
