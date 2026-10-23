"""Writes AlphaGoZero_Presentation_Explanations.docx: per-slide explanations, facts and Q&A.

Usage: python build_explanations.py <output.docx>
"""
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from explanations_content import EXPLANATIONS, GLOSSARY
from scripts_content import MEMBERS, SCRIPTS, SLIDE_TITLES

SLATE = RGBColor(0x5A, 0x64, 0x72)


def label(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)


def build(path):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11.5)

    doc.add_heading("AlphaGo Zero: Presentation Explanations", level=0)
    doc.add_paragraph("A study guide for the team. For each slide: the idea in plain language, the numbers to "
                      "remember, and questions the examiners may ask, with answers. Read the slides you present "
                      "closely, and skim the rest so anyone can answer a question on any slide.")

    for n in sorted(EXPLANATIONS):
        e = EXPLANATIONS[n]
        who, minutes, _ = SCRIPTS[n]
        doc.add_heading(f"Slide {n}: {SLIDE_TITLES[n]}", level=1)
        meta = doc.add_paragraph(f"Presented by {MEMBERS[who][0]} · about {minutes:g} min")
        meta.runs[0].italic = True
        meta.runs[0].font.color.rgb = SLATE

        label(doc, "In simple words")
        for t in e["simple"]:
            doc.add_paragraph(t)
        if e["remember"]:
            label(doc, "Remember")
            for t in e["remember"]:
                doc.add_paragraph(t, style="List Bullet")
        if e["qa"]:
            label(doc, "Possible questions")
            for q, a in e["qa"]:
                p = doc.add_paragraph()
                p.add_run("Q: ").bold = True
                p.add_run(q).bold = True
                p.paragraph_format.space_after = Pt(1)
                p = doc.add_paragraph()
                p.add_run("A: ").bold = True
                p.add_run(a)

    doc.add_page_break()
    doc.add_heading("Glossary", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Term", "Meaning"
    for term, meaning in GLOSSARY:
        row = table.add_row().cells
        row[0].text, row[1].text = term, meaning

    # python-docx's default template omits the required w:percent on w:zoom
    for zoom in doc.settings.element.iter(qn("w:zoom")):
        if zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")
    doc.save(path)


if __name__ == "__main__":
    build(sys.argv[1])
    print("wrote", sys.argv[1])
